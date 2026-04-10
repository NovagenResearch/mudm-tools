#!/usr/bin/env python3
"""Generate two PowerPoint slides showing the muDM JSON metadata pyramid."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Colours
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MID_GRAY = RGBColor(0x66, 0x66, 0x66)
LIGHT_BG = RGBColor(0xF5, 0xF5, 0xF5)

# Tier colours (muted blues/greens)
C_PYRAMID = RGBColor(0x2B, 0x57, 0x97)   # deep blue
C_TILE    = RGBColor(0x3A, 0x7C, 0xA5)   # teal
C_FEATURE = RGBColor(0x4C, 0xA1, 0x6C)   # green
C_DATA    = RGBColor(0xE8, 0x8D, 0x2A)   # amber

C_PYRAMID_FILL = RGBColor(0xD6, 0xE4, 0xF0)
C_TILE_FILL    = RGBColor(0xD4, 0xEC, 0xF0)
C_FEATURE_FILL = RGBColor(0xDA, 0xF0, 0xDB)
C_DATA_FILL    = RGBColor(0xFD, 0xF0, 0xD5)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height

# ── helpers ──────────────────────────────────────────────────────────
def add_text(slide, left, top, width, height, text, size=14,
             bold=False, color=DARK_GRAY, align=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = align
    return txBox


def add_rounded_rect(slide, left, top, width, height, fill_color, border_color, text="",
                     text_size=12, text_color=DARK_GRAY, bold=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    # Reduce rounding
    shape.adjustments[0] = 0.05
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(text_size)
        p.font.color.rgb = text_color
        p.font.bold = bold
        p.font.name = "Calibri"
    return shape


def add_arrow_down(slide, cx, top, length, color=MID_GRAY):
    """Draw a downward arrow (line + triangle)."""
    # Vertical line
    connector = slide.shapes.add_connector(
        1, cx, top, cx, top + length)  # 1 = straight connector
    connector.line.color.rgb = color
    connector.line.width = Pt(2)
    # Arrowhead via a small triangle
    tri_size = Inches(0.15)
    tri = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        cx - tri_size // 2, top + length - Pt(2),
        tri_size, tri_size)
    tri.fill.solid()
    tri.fill.fore_color.rgb = color
    tri.line.fill.background()
    tri.rotation = 180.0
    return connector


def add_code_block(slide, left, top, width, height, lines, title="",
                   title_color=DARK_GRAY):
    """Add a rounded rectangle with monospace text inside."""
    bg = add_rounded_rect(slide, left, top, width, height,
                          RGBColor(0xF8, 0xF8, 0xF8), RGBColor(0xCC, 0xCC, 0xCC))
    tf = bg.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.15)
    tf.margin_right = Inches(0.15)
    tf.margin_top = Inches(0.1)
    tf.margin_bottom = Inches(0.1)

    # Clear default paragraph
    tf.paragraphs[0].clear()

    if title:
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = title_color
        p.font.name = "Calibri"
        p.space_after = Pt(4)

    for i, line in enumerate(lines):
        p = tf.add_paragraph() if (i > 0 or title) else tf.paragraphs[0]
        p.text = line
        p.font.size = Pt(10)
        p.font.color.rgb = DARK_GRAY
        p.font.name = "Consolas"
        p.space_before = Pt(0)
        p.space_after = Pt(1)
    return bg


# =====================================================================
# SLIDE 1: Metadata Hierarchy Overview
# =====================================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
slide1.background.fill.solid()
slide1.background.fill.fore_color.rgb = WHITE

add_text(slide1, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6),
         "muDM Metadata Pyramid", size=28, bold=True, color=DARK_GRAY)
add_text(slide1, Inches(0.5), Inches(0.85), Inches(12), Inches(0.4),
         "Three-tier JSON hierarchy from dataset manifest to individual tile references",
         size=14, color=MID_GRAY)

# ── Tier boxes (left side, stacked vertically with arrows) ──
box_w = Inches(4.5)
box_h = Inches(1.25)
left_x = Inches(0.7)
y_start = Inches(1.7)
y_gap = Inches(1.85)

tiers = [
    ("pyramids.json", "Dataset Manifest",
     "Lists all available pyramids\nwith summary stats and paths",
     C_PYRAMID, C_PYRAMID_FILL),
    ("tilejson3d.json", "Pyramid Metadata",
     "Spatial bounds, zoom levels,\ntile encodings, vector layer schema",
     C_TILE, C_TILE_FILL),
    ("features.json", "Feature Index (MuDM)",
     "Per-feature properties and\ntile back-references for spatial lookup",
     C_FEATURE, C_FEATURE_FILL),
]

for i, (filename, label, desc, border_c, fill_c) in enumerate(tiers):
    y = y_start + i * y_gap

    box = add_rounded_rect(slide1, left_x, y, box_w, box_h, fill_c, border_c)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top = Inches(0.12)

    # Filename
    p = tf.paragraphs[0]
    p.text = filename
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = border_c
    p.font.name = "Consolas"

    # Label
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = DARK_GRAY
    p2.font.name = "Calibri"
    p2.space_before = Pt(2)

    # Description
    p3 = tf.add_paragraph()
    p3.text = desc
    p3.font.size = Pt(11)
    p3.font.color.rgb = MID_GRAY
    p3.font.name = "Calibri"
    p3.space_before = Pt(2)

    # Arrow between tiers
    if i < len(tiers) - 1:
        arrow_cx = left_x + box_w // 2
        arrow_top = y + box_h + Pt(2)
        arrow_len = y_gap - box_h - Pt(4)
        add_arrow_down(slide1, arrow_cx, arrow_top, arrow_len, border_c)

# ── Data tier (amber, below features) ──
data_y = y_start + 3 * y_gap
data_box = add_rounded_rect(
    slide1, left_x, data_y, box_w, Inches(0.8),
    C_DATA_FILL, C_DATA)
tf = data_box.text_frame
tf.word_wrap = True
tf.margin_left = Inches(0.2)
tf.margin_top = Inches(0.08)
p = tf.paragraphs[0]
p.text = "{z}/{x}/{y}/{d}.glb  |  .parquet"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = C_DATA
p.font.name = "Consolas"
p2 = tf.add_paragraph()
p2.text = "Binary tile data (GLB meshes, Parquet tables)"
p2.font.size = Pt(11)
p2.font.color.rgb = MID_GRAY
p2.font.name = "Calibri"
p2.space_before = Pt(2)

# Arrow from features to data
feat_y = y_start + 2 * y_gap
arrow_cx = left_x + box_w // 2
add_arrow_down(slide1, arrow_cx, feat_y + box_h + Pt(2),
               data_y - feat_y - box_h - Pt(4), C_FEATURE)

# ── Right side: directory tree ──
tree_left = Inches(5.8)
tree_top = Inches(1.7)

tree_lines = [
    "tiles-base/",
    "  pyramids.json",
    "  hemibrain_1k/",
    "    tilejson3d.json",
    "    features.json",
    "    3dtiles/",
    "      0/0/0/0.glb",
    "      1/1/1/0.glb",
    "      1/1/1/1.glb",
    "      ...",
    "    parquet/",
    "      0/0/0/0.parquet",
    "      ...",
    "  mouselight_500/",
    "    tilejson3d.json",
    "    features.json",
    "    3dtiles/",
    "      ...",
]

add_text(slide1, tree_left, tree_top - Inches(0.05), Inches(3.5), Inches(0.35),
         "Directory Layout", size=16, bold=True, color=DARK_GRAY)

tree_block = add_code_block(
    slide1, tree_left, tree_top + Inches(0.35), Inches(3.8), Inches(4.5),
    tree_lines)

# ── Relationship annotations on the right ──
ann_left = Inches(10.0)
annotations = [
    (Inches(2.2), "1 : N", "One manifest\nper tile store"),
    (Inches(3.5), "1 : 1", "One metadata file\nper pyramid"),
    (Inches(4.8), "1 : N", "Each feature lists\nits tile addresses"),
]
for y_off, ratio, note in annotations:
    add_text(slide1, ann_left, tree_top + y_off, Inches(1.0), Inches(0.35),
             ratio, size=20, bold=True, color=C_PYRAMID,
             align=PP_ALIGN.CENTER, font_name="Consolas")
    add_text(slide1, ann_left, tree_top + y_off + Inches(0.35), Inches(2.8), Inches(0.5),
             note, size=11, color=MID_GRAY, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 2: Data Flow — How tiles are referenced
# =====================================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
slide2.background.fill.solid()
slide2.background.fill.fore_color.rgb = WHITE

add_text(slide2, Inches(0.5), Inches(0.3), Inches(12), Inches(0.6),
         "Tile Data Resolution", size=28, bold=True, color=DARK_GRAY)
add_text(slide2, Inches(0.5), Inches(0.85), Inches(12), Inches(0.4),
         "How the viewer resolves a feature to its binary tile data across zoom levels",
         size=14, color=MID_GRAY)

# ── Left: features.json snippet ──
feat_left = Inches(0.5)
feat_top = Inches(1.6)
feat_w = Inches(4.2)

feat_lines = [
    '{',
    '  "type": "Feature",',
    '  "id": "ER5(ring)_L",',
    '  "geometry": {',
    '    "type": "TIN",',
    '    "coordinates": [],',
    '    "tiles": [',
    '      "0/0/0/0",',
    '      "1/1/1/0",',
    '      "2/2/2/1",',
    '      "3/5/5/3",',
    '      "4/10/10/7"',
    '    ]',
    '  },',
    '  "properties": {',
    '    "body_id": 1200057627,',
    '    "cell_type": "ER5",',
    '    "color": "#1fdfcf"',
    '  }',
    '}',
]

add_text(slide2, feat_left, feat_top - Inches(0.05), feat_w, Inches(0.35),
         "features.json  (per feature)", size=14, bold=True, color=C_FEATURE)
add_code_block(slide2, feat_left, feat_top + Inches(0.3), feat_w, Inches(4.6),
               feat_lines)

# ── Middle: tilejson3d.json snippet ──
tile_left = Inches(5.2)
tile_top = Inches(1.6)
tile_w = Inches(3.8)

tile_lines = [
    '{',
    '  "tilejson": "3.0.0",',
    '  "tiles": ["{z}/{x}/{y}/{d}"],',
    '  "minzoom": 0,',
    '  "maxzoom": 4,',
    '  "bounds3d": [',
    '    16672, 79696, 34560,',
    '    275424, 266672, 280496',
    '  ],',
    '  "encodings": [{',
    '    "format": "glb",',
    '    "compression": "meshopt",',
    '    "path": "3dtiles",',
    '    "extension": ".glb"',
    '  }],',
    '  "zoom_counts": {',
    '    "0": 1, "1": 8,',
    '    "2": 33, "3": 108,',
    '    "4": 346',
    '  }',
    '}',
]

add_text(slide2, tile_left, tile_top - Inches(0.05), tile_w, Inches(0.35),
         "tilejson3d.json  (per pyramid)", size=14, bold=True, color=C_TILE)
add_code_block(slide2, tile_left, tile_top + Inches(0.3), tile_w, Inches(4.6),
               tile_lines)

# ── Right: Resolution logic ──
res_left = Inches(9.5)
res_top = Inches(1.6)
res_w = Inches(3.5)

add_text(slide2, res_left, res_top - Inches(0.05), res_w, Inches(0.35),
         "Tile Resolution", size=14, bold=True, color=C_DATA)

steps = [
    ("1", "Select feature", 'Viewer picks feature\nby id or spatial query'),
    ("2", "Read tile list", 'geometry.tiles gives\nall tile addresses'),
    ("3", "Pick zoom level", 'Filter tiles matching\ncurrent zoom level z'),
    ("4", "Resolve file path", '{path}/{z}/{x}/{y}/{d}{ext}\n3dtiles/4/10/10/7.glb'),
    ("5", "Fetch & render", 'Load binary tile,\ndecode, display'),
]

step_h = Inches(0.75)
step_gap = Inches(0.15)
for i, (num, title, desc) in enumerate(steps):
    sy = res_top + Inches(0.35) + i * (step_h + step_gap)

    # Number circle
    circle = slide2.shapes.add_shape(
        MSO_SHAPE.OVAL,
        res_left, sy + Pt(4), Inches(0.35), Inches(0.35))
    circle.fill.solid()
    circle.fill.fore_color.rgb = C_TILE
    circle.line.fill.background()
    tf = circle.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Calibri"

    # Step text
    add_text(slide2, res_left + Inches(0.45), sy,
             res_w - Inches(0.5), Inches(0.25),
             title, size=12, bold=True, color=DARK_GRAY)
    add_text(slide2, res_left + Inches(0.45), sy + Inches(0.22),
             res_w - Inches(0.5), Inches(0.55),
             desc, size=10, color=MID_GRAY)

# ── Connecting arrows between code blocks ──
# Arrow from features tiles field to tilejson
arrow_y = feat_top + Inches(2.8)
connector = slide2.shapes.add_connector(
    1, feat_left + feat_w + Pt(4), arrow_y,
    tile_left - Pt(4), arrow_y)
connector.line.color.rgb = C_FEATURE
connector.line.width = Pt(2)

# Label on arrow
add_text(slide2, feat_left + feat_w + Inches(0.05), arrow_y - Inches(0.3),
         Inches(0.8), Inches(0.25),
         "tile IDs", size=10, bold=True, color=C_FEATURE,
         align=PP_ALIGN.CENTER)

# Arrow from tilejson encodings to resolution steps
arrow_y2 = tile_top + Inches(3.2)
connector2 = slide2.shapes.add_connector(
    1, tile_left + tile_w + Pt(4), arrow_y2,
    res_left - Pt(4), arrow_y2)
connector2.line.color.rgb = C_TILE
connector2.line.width = Pt(2)

add_text(slide2, tile_left + tile_w + Inches(0.05), arrow_y2 - Inches(0.3),
         Inches(0.8), Inches(0.25),
         "encoding", size=10, bold=True, color=C_TILE,
         align=PP_ALIGN.CENTER)

# ── Save ──
out_path = "/data/code/work/microjson/muDM_metadata_pyramid.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
