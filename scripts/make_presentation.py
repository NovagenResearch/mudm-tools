#!/usr/bin/env python3
"""Generate the muDM 17-slide presentation as PowerPoint."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── Palette ──────────────────────────────────────────────────────────
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BLACK      = RGBColor(0x00, 0x00, 0x00)
DARK       = RGBColor(0x2D, 0x2D, 0x2D)
MID        = RGBColor(0x66, 0x66, 0x66)
LIGHT      = RGBColor(0x99, 0x99, 0x99)
ACCENT1    = RGBColor(0x2B, 0x57, 0x97)  # deep blue
ACCENT2    = RGBColor(0x3A, 0x7C, 0xA5)  # teal
ACCENT3    = RGBColor(0x4C, 0xA1, 0x6C)  # green
ACCENT4    = RGBColor(0xE8, 0x8D, 0x2A)  # amber
ACCENT5    = RGBColor(0x8B, 0x5C, 0xF6)  # purple
BG_LIGHT   = RGBColor(0xF7, 0xF8, 0xFA)
BG_BLUE    = RGBColor(0xE8, 0xEF, 0xF7)
BG_TEAL    = RGBColor(0xE4, 0xF0, 0xF3)
BG_GREEN   = RGBColor(0xE6, 0xF4, 0xE8)
BG_AMBER   = RGBColor(0xFD, 0xF0, 0xD5)
BG_PURPLE  = RGBColor(0xEE, 0xE8, 0xFD)
CODE_BG    = RGBColor(0xF0, 0xF0, 0xF0)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW = prs.slide_width
SH = prs.slide_height


# ── Helpers ──────────────────────────────────────────────────────────
def _blank(bg_color=WHITE):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = bg_color
    return sl


def _txt(sl, left, top, w, h, text, sz=14, bold=False, color=DARK,
         align=PP_ALIGN.LEFT, font="Calibri", italic=False):
    tb = sl.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    p.font.italic = italic
    p.alignment = align
    return tb


def _multi(sl, left, top, w, h, lines, sz=14, color=DARK, font="Calibri",
           bold=False, spacing=Pt(6), align=PP_ALIGN.LEFT, bullet=False):
    """Multi-line text box. lines is list of str."""
    tb = sl.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(sz)
        p.font.color.rgb = color
        p.font.name = font
        p.font.bold = bold
        p.space_before = spacing if i > 0 else Pt(0)
        p.alignment = align
        if bullet and i > 0:
            p.level = 0
    return tb


def _rect(sl, l, t, w, h, fill, border=None, radius=0.08):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if border:
        s.line.color.rgb = border
        s.line.width = Pt(1.5)
    else:
        s.line.fill.background()
    s.adjustments[0] = radius
    return s


def _circle(sl, l, t, size, fill):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, l, t, size, size)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def _arrow_right(sl, x1, y1, x2, y2=None, color=MID, width=2):
    if y2 is None:
        y2 = y1
    c = sl.shapes.add_connector(1, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    # arrowhead triangle
    tri_sz = Inches(0.18)
    tri = sl.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                              x2 - Pt(2), y2 - tri_sz // 2, tri_sz, tri_sz)
    tri.fill.solid()
    tri.fill.fore_color.rgb = color
    tri.line.fill.background()
    tri.rotation = 90.0
    return c


def _arrow_down(sl, cx, y1, y2, color=MID, width=2):
    c = sl.shapes.add_connector(1, cx, y1, cx, y2)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    tri_sz = Inches(0.15)
    tri = sl.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                              cx - tri_sz // 2, y2 - Pt(2), tri_sz, tri_sz)
    tri.fill.solid()
    tri.fill.fore_color.rgb = color
    tri.line.fill.background()
    tri.rotation = 180.0
    return c


def _section_tag(sl, text, color=ACCENT1):
    """Small section label in top-left."""
    _txt(sl, Inches(0.5), Inches(0.25), Inches(4), Inches(0.3),
         text, sz=11, color=color, bold=True)


def _slide_title(sl, text, subtitle=None, tag=None, tag_color=ACCENT1):
    if tag:
        _section_tag(sl, tag, tag_color)
    _txt(sl, Inches(0.5), Inches(0.6), Inches(12), Inches(0.7),
         text, sz=32, bold=True, color=DARK)
    if subtitle:
        _txt(sl, Inches(0.5), Inches(1.25), Inches(11), Inches(0.5),
             subtitle, sz=16, color=MID)


def _figure_placeholder(sl, left, top, w, h, label):
    """Gray box with figure prompt label."""
    r = _rect(sl, left, top, w, h, BG_LIGHT, RGBColor(0xCC, 0xCC, 0xCC))
    tf = r.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(12)
    p.font.color.rgb = LIGHT
    p.font.italic = True
    p.font.name = "Calibri"
    return r


def _icon_box(sl, left, top, w, h, label, fill, border, label_sz=13):
    """Rounded rect with centered label."""
    r = _rect(sl, left, top, w, h, fill, border)
    tf = r.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(label_sz)
    p.font.bold = True
    p.font.color.rgb = border
    p.font.name = "Calibri"
    return r


def _big_number(sl, left, top, number, label, color=ACCENT1, sub_label=None):
    """Large metric number with label below."""
    _txt(sl, left, top, Inches(3.5), Inches(0.9),
         number, sz=48, bold=True, color=color, align=PP_ALIGN.CENTER,
         font="Calibri")
    _txt(sl, left, top + Inches(0.85), Inches(3.5), Inches(0.5),
         label, sz=14, color=DARK, align=PP_ALIGN.CENTER)
    if sub_label:
        _txt(sl, left, top + Inches(1.25), Inches(3.5), Inches(0.35),
             sub_label, sz=11, color=LIGHT, align=PP_ALIGN.CENTER,
             italic=True)


# =====================================================================
# SLIDE 1 — Title
# =====================================================================
sl = _blank(WHITE)
_txt(sl, Inches(1.5), Inches(2.0), Inches(10), Inches(1.2),
     "muDM", sz=64, bold=True, color=ACCENT1, align=PP_ALIGN.CENTER)
_txt(sl, Inches(1.5), Inches(3.2), Inches(10), Inches(0.8),
     "A Multi-Format Spatial Data Model\nfor ML-Native Bioimaging",
     sz=24, color=DARK, align=PP_ALIGN.CENTER)
_txt(sl, Inches(1.5), Inches(5.0), Inches(10), Inches(0.4),
     "Bengt Ljungquist", sz=18, bold=True, color=MID, align=PP_ALIGN.CENTER)
_txt(sl, Inches(1.5), Inches(5.5), Inches(10), Inches(0.4),
     "Novagen Research Fund / Nextonic Solutions", sz=14, color=LIGHT,
     align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 2 — The Problem
# =====================================================================
sl = _blank()
_slide_title(sl, "The Problem", tag="MOTIVATION")

# Format silo icons
formats = ["OBJ", "GeoJSON", "Zarr\nlabels", "Proprietary\nexports"]
fmt_colors = [ACCENT1, ACCENT2, ACCENT3, ACCENT4]
fmt_bg     = [BG_BLUE, BG_TEAL, BG_GREEN, BG_AMBER]
box_w = Inches(1.6)
box_h = Inches(1.0)
start_x = Inches(0.8)
start_y = Inches(2.5)
gap = Inches(0.4)

for i, (fmt, fc, bg) in enumerate(zip(formats, fmt_colors, fmt_bg)):
    x = start_x + i * (box_w + gap)
    _icon_box(sl, x, start_y, box_w, box_h, fmt, bg, fc)

# Crossed arrows between them
for i in range(len(formats) - 1):
    x1 = start_x + (i + 1) * (box_w + gap) - gap + Pt(4)
    x2 = start_x + (i + 1) * (box_w + gap) - Pt(4)
    # just a red X between boxes
    cx = start_x + i * (box_w + gap) + box_w + gap // 2
    _txt(sl, cx - Inches(0.15), start_y + Inches(0.25), Inches(0.3), Inches(0.5),
         "X", sz=20, bold=True, color=RGBColor(0xCC, 0x33, 0x33),
         align=PP_ALIGN.CENTER)

# Problem bullets on the right
bullets = [
    "Visualization tools cannot read ML formats",
    "ML pipelines cannot read visualization formats",
    "Every new tool = another conversion script",
    "Metadata is lost at every conversion step",
]
for i, b in enumerate(bullets):
    y = Inches(2.3) + i * Inches(0.55)
    _txt(sl, Inches(8.5), y, Inches(4.5), Inches(0.5),
         b, sz=16, color=DARK)
    # bullet dot
    _circle(sl, Inches(8.15), y + Pt(5), Inches(0.12), RGBColor(0xCC, 0x33, 0x33))

# Bottom message
_txt(sl, Inches(0.8), Inches(5.5), Inches(11), Inches(0.5),
     "Bioimaging produces rich spatial annotations, but they are trapped in format silos.",
     sz=18, bold=True, color=ACCENT1, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 3 — What if one data model served all three?
# =====================================================================
sl = _blank()
_slide_title(sl, "What if one data model served all three?", tag="VISION")

# Central muDM node
cx, cy = Inches(6.666), Inches(4.0)
mudm_sz = Inches(2.0)
c = _circle(sl, cx - mudm_sz // 2, cy - mudm_sz // 2, mudm_sz, ACCENT1)
tf = c.text_frame
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.text = "muDM"
p.font.size = Pt(28)
p.font.bold = True
p.font.color.rgb = WHITE
p.font.name = "Calibri"

# Three targets
targets = [
    (Inches(2.0), Inches(2.5), "Web\nVisualization", ACCENT2, BG_TEAL),
    (Inches(10.0), Inches(2.5), "ML\nTraining", ACCENT3, BG_GREEN),
    (Inches(6.0), Inches(6.0), "Interoperability\n& Analysis", ACCENT4, BG_AMBER),
]
for tx, ty, label, tc, tbg in targets:
    _icon_box(sl, tx, ty, Inches(2.2), Inches(1.0), label, tbg, tc, label_sz=15)

# Lines from center to targets (simple connectors)
for tx, ty, _, tc, _ in targets:
    target_cx = tx + Inches(1.1)
    target_cy = ty + Inches(0.5)
    c = sl.shapes.add_connector(1, cx, cy, target_cx, target_cy)
    c.line.color.rgb = tc
    c.line.width = Pt(2.5)

# Subtitle
_txt(sl, Inches(1.0), Inches(1.5), Inches(11), Inches(0.4),
     "One annotation model. No conversion. Ontology-linked metadata travels with the geometry.",
     sz=16, color=MID, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 4 — From MuDM to muDM
# =====================================================================
sl = _blank()
_slide_title(sl, "From MuDM to muDM", tag="HISTORY")

# Left: MuDM (small)
mj_x, mj_y = Inches(0.8), Inches(2.8)
mj_w, mj_h = Inches(3.0), Inches(2.5)
r = _rect(sl, mj_x, mj_y, mj_w, mj_h, BG_LIGHT, ACCENT2)
_txt(sl, mj_x + Inches(0.2), mj_y + Inches(0.15), Inches(2.6), Inches(0.4),
     "MuDM", sz=20, bold=True, color=ACCENT2, font="Consolas")
_txt(sl, mj_x + Inches(0.2), mj_y + Inches(0.6), Inches(2.6), Inches(0.4),
     "GeoJSON-based format", sz=13, color=MID)
_icon_box(sl, mj_x + Inches(0.5), mj_y + Inches(1.2), Inches(2.0), Inches(0.7),
          "features.json", BG_TEAL, ACCENT2, label_sz=12)
_txt(sl, mj_x + Inches(0.2), mj_y + Inches(2.05), Inches(2.6), Inches(0.4),
     "Single file format", sz=12, color=LIGHT, italic=True)

# Arrow
_arrow_right(sl, mj_x + mj_w + Inches(0.2), Inches(4.0),
             Inches(5.3), Inches(4.0), ACCENT1, 3)

# Right: muDM (expanded)
md_x, md_y = Inches(5.5), Inches(2.3)
md_w, md_h = Inches(7.0), Inches(3.8)
r = _rect(sl, md_x, md_y, md_w, md_h, BG_BLUE, ACCENT1)
_txt(sl, md_x + Inches(0.3), md_y + Inches(0.15), Inches(3), Inches(0.5),
     "muDM", sz=24, bold=True, color=ACCENT1, font="Consolas")
_txt(sl, md_x + Inches(2.0), md_y + Inches(0.2), Inches(4.5), Inches(0.35),
     "micro Data Model", sz=14, color=MID)

# Components inside muDM box
components = [
    ("Data Model", ACCENT1, BG_LIGHT),
    ("Tiling Pipeline", ACCENT2, BG_TEAL),
    ("Multi-Format Output", ACCENT3, BG_GREEN),
    ("Web Viewer", ACCENT4, BG_AMBER),
]
comp_w = Inches(1.45)
comp_h = Inches(0.65)
for i, (label, cc, cbg) in enumerate(components):
    cx = md_x + Inches(0.3) + i * (comp_w + Inches(0.1))
    _icon_box(sl, cx, md_y + Inches(0.8), comp_w, comp_h, label, cbg, cc, 11)

# MuDM subset indicator inside muDM
_icon_box(sl, md_x + Inches(0.3), md_y + Inches(1.8), Inches(2.5), Inches(0.6),
          "MuDM = features.json", BG_TEAL, ACCENT2, 11)

_txt(sl, md_x + Inches(0.3), md_y + Inches(2.6), Inches(6.4), Inches(0.8),
     "Ontology support  |  Streaming generation  |  Open source",
     sz=13, color=MID)

# Bottom tagline
_txt(sl, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.5),
     "Not just a format \u2014 a data model.",
     sz=20, bold=True, color=ACCENT1, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 5 — muDM at a glance
# =====================================================================
sl = _blank()
_slide_title(sl, "muDM at a Glance", tag="ARCHITECTURE")

# Top: MuDM Data Model box
top_x, top_y = Inches(3.5), Inches(1.8)
top_w = Inches(6.0)
_icon_box(sl, top_x, top_y, top_w, Inches(0.8),
          "MuDM Data Model (GeoJSON-inspired)", BG_BLUE, ACCENT1, 16)

# Arrow down
_arrow_down(sl, top_x + top_w // 2, top_y + Inches(0.8) + Pt(4),
            Inches(3.2), ACCENT1)

# Middle row: Python API + Rust Engine
_icon_box(sl, Inches(2.5), Inches(3.4), Inches(3.5), Inches(0.8),
          "Python API", BG_TEAL, ACCENT2, 16)
_icon_box(sl, Inches(7.0), Inches(3.4), Inches(3.5), Inches(0.8),
          "Rust Tiling Engine (PyO3)", BG_GREEN, ACCENT3, 16)

# Arrows down from both
_arrow_down(sl, Inches(4.25), Inches(4.2) + Pt(4), Inches(4.8), ACCENT2)
_arrow_down(sl, Inches(8.75), Inches(4.2) + Pt(4), Inches(4.8), ACCENT3)

# Bottom row: output formats
outputs = [
    ("OGC 3D Tiles\n(GLB)", ACCENT1, BG_BLUE),
    ("Apache Parquet\n(ZSTD)", ACCENT3, BG_GREEN),
    ("Vector Tiles\n(MVT)", ACCENT2, BG_TEAL),
    ("Neuroglancer", ACCENT5, BG_PURPLE),
]
out_w = Inches(2.4)
out_start = Inches(1.5)
out_gap = Inches(0.4)
for i, (label, oc, obg) in enumerate(outputs):
    _icon_box(sl, out_start + i * (out_w + out_gap), Inches(5.2),
              out_w, Inches(0.9), label, obg, oc, 13)

# Ontology badge
_icon_box(sl, Inches(10.5), Inches(1.8), Inches(2.3), Inches(0.8),
          "Ontology-linked\nmetadata", BG_AMBER, ACCENT4, 12)
_txt(sl, Inches(10.5), Inches(6.5), Inches(2.3), Inches(0.4),
     "Open source", sz=14, bold=True, color=ACCENT3, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 6 — Use cases
# =====================================================================
sl = _blank()
_slide_title(sl, "Use Cases", tag="APPLICATIONS")

cases = [
    ("Connectomics", "3D neuron meshes across\nwhole-brain atlases",
     ACCENT1, BG_BLUE),
    ("Spatial Omics / Oncology", "Cell-level annotations\nover tissue sections",
     ACCENT3, BG_GREEN),
    ("General Bioimaging", "ROIs, segmentation masks,\ntracked objects",
     ACCENT4, BG_AMBER),
]
panel_w = Inches(3.5)
panel_h = Inches(3.5)
panel_gap = Inches(0.5)
start_x = Inches(1.2)

for i, (title, desc, cc, cbg) in enumerate(cases):
    px = start_x + i * (panel_w + panel_gap)
    py = Inches(2.2)

    r = _rect(sl, px, py, panel_w, panel_h, cbg, cc)

    # Placeholder for illustration
    _figure_placeholder(sl, px + Inches(0.3), py + Inches(0.3),
                        panel_w - Inches(0.6), Inches(1.6),
                        "[illustration]")

    _txt(sl, px + Inches(0.3), py + Inches(2.1), panel_w - Inches(0.6), Inches(0.4),
         title, sz=18, bold=True, color=cc)
    _txt(sl, px + Inches(0.3), py + Inches(2.55), panel_w - Inches(0.6), Inches(0.8),
         desc, sz=14, color=MID)

_txt(sl, Inches(1.2), Inches(6.2), Inches(10), Inches(0.4),
     "Same data model handles all three.", sz=18, bold=True, color=ACCENT1,
     align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 7 — Vector vs. raster
# =====================================================================
sl = _blank()
_slide_title(sl, "Vector vs. Raster for Annotations", tag="APPROACH")

# Left: Raster
rx, ry = Inches(0.8), Inches(2.2)
rw, rh = Inches(5.5), Inches(4.0)
r = _rect(sl, rx, ry, rw, rh, BG_LIGHT, LIGHT)
_txt(sl, rx + Inches(0.3), ry + Inches(0.15), Inches(4), Inches(0.4),
     "Raster (e.g. OME-Zarr labels)", sz=18, bold=True, color=MID)

# Pixel grid placeholder
_figure_placeholder(sl, rx + Inches(0.3), ry + Inches(0.7),
                    rw - Inches(0.6), Inches(1.8), "[colored pixel grid]")

_multi(sl, rx + Inches(0.3), ry + Inches(2.7), rw - Inches(0.6), Inches(1.2),
       ["Dense pixel masks", "No per-object identity",
        "No inline metadata", "Good for: segmentation masks, label volumes"],
       sz=14, color=MID, spacing=Pt(4))

# Right: Vector
vx = Inches(7.0)
r = _rect(sl, vx, ry, rw, rh, BG_GREEN, ACCENT3)
_txt(sl, vx + Inches(0.3), ry + Inches(0.15), Inches(4), Inches(0.4),
     "Vector (muDM)", sz=18, bold=True, color=ACCENT3)

_figure_placeholder(sl, vx + Inches(0.3), ry + Inches(0.7),
                    rw - Inches(0.6), Inches(1.8),
                    "[outlined objects with metadata tags]")

_multi(sl, vx + Inches(0.3), ry + Inches(2.7), rw - Inches(0.6), Inches(1.2),
       ["Discrete objects with identity", "Per-object metadata & ontology links",
        "Filterable and queryable", "Good for: annotations, ROIs, meshes"],
       sz=14, color=DARK, spacing=Pt(4))

# "vs" in center
_txt(sl, Inches(6.1), Inches(3.8), Inches(0.8), Inches(0.5),
     "vs.", sz=20, bold=True, color=LIGHT, align=PP_ALIGN.CENTER)

_txt(sl, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.4),
     "Not either/or \u2014 they complement each other.",
     sz=18, bold=True, color=ACCENT1, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 8 — One model, many formats
# =====================================================================
sl = _blank()
_slide_title(sl, "One Model, Many Formats", tag="PIPELINE")

# Large figure placeholder for Paper Fig. 1
_figure_placeholder(sl, Inches(0.8), Inches(1.8), Inches(11.5), Inches(4.5),
                    "[INSERT: Paper Fig. 1 \u2014 fig1_pipeline.pdf]\n\n"
                    "muDM multi-format tiling pipeline: spatial indexing (octree/quadtree)\n"
                    "\u2192 parallel encoding into OGC 3D Tiles, Vector Tiles, Apache Parquet")

_txt(sl, Inches(0.8), Inches(6.5), Inches(11.5), Inches(0.5),
     "Also supports: Neuroglancer precomputed meshes  |  Standard MVT vector tiles (Mapbox/OpenLayers)",
     sz=14, color=MID, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 9 — OGC 3D Tiles
# =====================================================================
sl = _blank()
_slide_title(sl, "OGC 3D Tiles: Web Visualization", tag="FORMAT")

# Left: text
bullets = [
    "Open standard from OGC",
    "GLB meshes with Meshopt compression",
    "Level-of-detail via octree",
    "Browser loads only visible tiles at current zoom",
    "Metadata embedded per node for click-to-inspect",
    "Compatible with Three.js / CesiumJS viewers",
]
for i, b in enumerate(bullets):
    y = Inches(2.0) + i * Inches(0.55)
    _circle(sl, Inches(0.8), y + Pt(6), Inches(0.12), ACCENT1)
    _txt(sl, Inches(1.1), y, Inches(5.0), Inches(0.5), b, sz=16, color=DARK)

# Right: LOD diagram placeholder
_figure_placeholder(sl, Inches(6.8), Inches(1.8), Inches(5.8), Inches(4.5),
                    "[LOD diagram]\n\n"
                    "z=0: coarse mesh\n"
                    "z=2: medium detail\n"
                    "z=4: full detail\n\n"
                    "Same object at increasing fidelity")


# =====================================================================
# SLIDE 10 — Apache Parquet: ML training
# =====================================================================
sl = _blank()
_slide_title(sl, "Apache Parquet: ML Training", tag="FORMAT")

# Top path (traditional) - grayed out
path_y1 = Inches(2.5)
trad_steps = ["Raw mesh\nfiles", "Parse", "Convert", "Load", "PyTorch\nDataLoader"]
trad_w = Inches(1.8)
trad_gap = Inches(0.5)
trad_start = Inches(1.0)

for i, step in enumerate(trad_steps):
    sx = trad_start + i * (trad_w + trad_gap)
    _icon_box(sl, sx, path_y1, trad_w, Inches(0.8), step,
              RGBColor(0xF0, 0xF0, 0xF0), LIGHT, 12)
    if i < len(trad_steps) - 1:
        _txt(sl, sx + trad_w + Inches(0.1), path_y1 + Inches(0.2),
             Inches(0.3), Inches(0.4), "\u2192", sz=20, color=LIGHT,
             align=PP_ALIGN.CENTER)

_txt(sl, Inches(0.3), path_y1 + Inches(0.2), Inches(0.7), Inches(0.4),
     "Traditional", sz=11, color=LIGHT, italic=True)

# Bottom path (Parquet) - highlighted
path_y2 = Inches(4.2)
_icon_box(sl, Inches(1.0), path_y2, Inches(2.5), Inches(1.0),
          "muDM Parquet\n(ZSTD compressed)", BG_GREEN, ACCENT3, 14)

_arrow_right(sl, Inches(3.6), path_y2 + Inches(0.5),
             Inches(7.5), path_y2 + Inches(0.5), ACCENT3, 3)

_icon_box(sl, Inches(7.7), path_y2, Inches(2.5), Inches(1.0),
          "PyTorch\nDataLoader", BG_GREEN, ACCENT3, 14)

_txt(sl, Inches(0.3), path_y2 + Inches(0.3), Inches(0.7), Inches(0.4),
     "muDM", sz=11, color=ACCENT3, bold=True)

# Benefits
benefits = [
    "Zero-copy tensor loading \u2014 no decode step",
    "Geometry + metadata in the same row",
    "Queryable: DuckDB, Pandas, PyArrow, Polars",
    "Spatial and metadata predicate pushdown",
]
for i, b in enumerate(benefits):
    y = Inches(5.7) + i * Inches(0.4)
    _circle(sl, Inches(1.0), y + Pt(6), Inches(0.1), ACCENT3)
    _txt(sl, Inches(1.3), y, Inches(10), Inches(0.35), b, sz=14, color=DARK)


# =====================================================================
# SLIDE 11 — OME-Zarr: complement, not competitor
# =====================================================================
sl = _blank()
_slide_title(sl, "OME-Zarr: Complement, Not Competitor", tag="ECOSYSTEM")

# Bottom layer: OME-Zarr
layer_x, layer_w = Inches(2.0), Inches(9.0)
zarr_y = Inches(4.2)
zarr_h = Inches(1.5)
r = _rect(sl, layer_x, zarr_y, layer_w, zarr_h, BG_LIGHT, LIGHT)
_txt(sl, layer_x + Inches(0.3), zarr_y + Inches(0.15), Inches(3), Inches(0.4),
     "OME-Zarr", sz=20, bold=True, color=MID)
_txt(sl, layer_x + Inches(0.3), zarr_y + Inches(0.6), Inches(8), Inches(0.7),
     "Raster image data: pixels, voxels, intensity volumes\n"
     "The emerging standard for bioimaging imagery",
     sz=14, color=MID)

# Top layer: muDM
mudm_y = Inches(2.2)
mudm_h = Inches(1.5)
r = _rect(sl, layer_x, mudm_y, layer_w, mudm_h, BG_GREEN, ACCENT3)
_txt(sl, layer_x + Inches(0.3), mudm_y + Inches(0.15), Inches(3), Inches(0.4),
     "muDM", sz=20, bold=True, color=ACCENT3)
_txt(sl, layer_x + Inches(0.3), mudm_y + Inches(0.6), Inches(8), Inches(0.7),
     "Vector annotation layer: objects, boundaries, metadata\n"
     "Discrete features with identity and ontology links",
     sz=14, color=DARK)

# Shared coordinate system indicator
_txt(sl, layer_x + layer_w + Inches(0.3), Inches(3.2), Inches(1.8), Inches(1.0),
     "Shared\ncoordinate\nsystem", sz=14, bold=True, color=ACCENT1,
     align=PP_ALIGN.CENTER)

# Brace/arrow between layers
_txt(sl, layer_x + layer_w + Inches(0.5), Inches(3.0), Inches(0.5), Inches(1.5),
     "}", sz=48, color=ACCENT1, align=PP_ALIGN.CENTER)

_txt(sl, Inches(2.0), Inches(6.2), Inches(9.0), Inches(0.5),
     "OME-Zarr for the image. muDM for what is in the image.",
     sz=18, bold=True, color=ACCENT1, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 12 — Streaming tile generation
# =====================================================================
sl = _blank()
_slide_title(sl, "Streaming Tile Generation", tag="SCALING")

# Left: large dataset
ds_x, ds_y = Inches(0.5), Inches(2.5)
_icon_box(sl, ds_x, ds_y, Inches(2.2), Inches(1.5),
          "> 1 TB\ndataset", BG_AMBER, ACCENT4, 18)

# Arrow to funnel
_arrow_right(sl, ds_x + Inches(2.3), ds_y + Inches(0.75),
             Inches(3.5), ds_y + Inches(0.75), ACCENT4, 3)

# Buckets
bucket_start = Inches(3.7)
bucket_w = Inches(0.9)
bucket_gap = Inches(0.15)
n_buckets = 5
for i in range(n_buckets):
    bx = bucket_start + i * (bucket_w + bucket_gap)
    _icon_box(sl, bx, ds_y + Inches(0.15), bucket_w, Inches(1.2),
              f"B{i+1}", BG_TEAL, ACCENT2, 11)

_txt(sl, bucket_start, ds_y - Inches(0.4), Inches(5.5), Inches(0.35),
     "Spatial buckets (processed independently)", sz=13, color=MID,
     align=PP_ALIGN.CENTER)

# Arrows from buckets to output
_arrow_right(sl, bucket_start + n_buckets * (bucket_w + bucket_gap),
             ds_y + Inches(0.75),
             Inches(9.5), ds_y + Inches(0.75), ACCENT2, 3)

# Output: tile grid
_icon_box(sl, Inches(9.7), ds_y, Inches(3.0), Inches(1.5),
          "Tile\nPyramid", BG_GREEN, ACCENT3, 18)

# Key points
points = [
    "Datasets can exceed RAM (thousands of neurons, millions of cells)",
    "Fragments spatially sorted into buckets, each processed independently",
    "Single workstation \u2014 no cluster required",
    "Constant memory, linear time",
]
for i, pt in enumerate(points):
    y = Inches(4.8) + i * Inches(0.5)
    _circle(sl, Inches(1.5), y + Pt(6), Inches(0.1), ACCENT2)
    _txt(sl, Inches(1.8), y, Inches(10), Inches(0.4), pt, sz=15, color=DARK)


# =====================================================================
# SLIDE 13 — Metadata pyramid (placeholder for reuse)
# =====================================================================
sl = _blank()
_slide_title(sl, "Metadata Pyramid", tag="DATA MODEL")

_figure_placeholder(sl, Inches(0.8), Inches(1.8), Inches(11.5), Inches(4.8),
                    "[INSERT: Slides from muDM_metadata_pyramid.pptx]\n\n"
                    "Three-tier JSON hierarchy:\n"
                    "pyramids.json (manifest)\n"
                    "\u2192 tilejson3d.json (spatial bounds, zoom levels, encodings)\n"
                    "\u2192 features.json (per-feature properties + tile back-references)\n\n"
                    "Or replace this slide with the two dedicated metadata pyramid slides")


# =====================================================================
# SLIDE 14 — Demo
# =====================================================================
sl = _blank(RGBColor(0x1A, 0x1A, 0x2E))
_txt(sl, Inches(2.0), Inches(2.0), Inches(9), Inches(1.5),
     "Demo", sz=72, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

_txt(sl, Inches(2.0), Inches(4.0), Inches(9), Inches(0.5),
     "Two things to watch for:", sz=18, color=RGBColor(0xAA, 0xAA, 0xCC),
     align=PP_ALIGN.CENTER)

_txt(sl, Inches(2.0), Inches(4.7), Inches(9), Inches(0.5),
     "1.  Progressive tile loading \u2014 tiles stream in as you zoom, coarse to fine",
     sz=16, color=RGBColor(0xDD, 0xDD, 0xEE), align=PP_ALIGN.CENTER)
_txt(sl, Inches(2.0), Inches(5.3), Inches(9), Inches(0.5),
     "2.  Metadata-driven filtering \u2014 select by cell type, region, or any property",
     sz=16, color=RGBColor(0xDD, 0xDD, 0xEE), align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 15 — Key results
# =====================================================================
sl = _blank()
_slide_title(sl, "Key Results", tag="PERFORMANCE")

# Top row: 3 metrics
_big_number(sl, Inches(0.5), Inches(2.0),
            "> 1 TB", "tiled in under 8 hours",
            ACCENT1, "single workstation, no cluster")

_big_number(sl, Inches(4.8), Inches(2.0),
            "21x", "faster full dataset loading",
            ACCENT3, "Parquet vs raw mesh files")

_big_number(sl, Inches(9.1), Inches(2.0),
            "~100\u20131000x", "faster selective queries",
            ACCENT5, "spatial/metadata filtering (est.)")

# Bottom row: 2 metrics
_big_number(sl, Inches(2.5), Inches(4.5),
            "141x", "faster decode",
            ACCENT2, "flat binary vs scene-graph parsing")

_big_number(sl, Inches(7.5), Inches(4.5),
            "~6x", "compression",
            ACCENT4, "ZSTD")

_txt(sl, Inches(0.5), Inches(6.7), Inches(12), Inches(0.4),
     "Same ML accuracy either way (95\u201396% on classification benchmarks).",
     sz=14, color=MID, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 16 — What's next
# =====================================================================
sl = _blank()
_slide_title(sl, "What\u2019s Next", tag="FUTURE")

# Timeline nodes
nodes = [
    ("Paper", "Spatial omics in oncology\nmuDM as annotation layer",
     ACCENT1, BG_BLUE),
    ("GSK Collaboration", "muDM as a data model\nJane leading",
     ACCENT3, BG_GREEN),
    ("Grant with GMU", "Expanding the platform\nand data model",
     ACCENT4, BG_AMBER),
]

node_w = Inches(3.2)
node_h = Inches(2.0)
node_gap = Inches(0.7)
start_x = Inches(0.8)
node_y = Inches(2.8)

for i, (title, desc, nc, nbg) in enumerate(nodes):
    nx = start_x + i * (node_w + node_gap)

    # Circle number
    c = _circle(sl, nx + node_w // 2 - Inches(0.25), node_y - Inches(0.5),
                Inches(0.5), nc)
    tf = c.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = str(i + 1)
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = WHITE

    r = _rect(sl, nx, node_y, node_w, node_h, nbg, nc)
    _txt(sl, nx + Inches(0.2), node_y + Inches(0.2), node_w - Inches(0.4), Inches(0.4),
         title, sz=18, bold=True, color=nc)
    _txt(sl, nx + Inches(0.2), node_y + Inches(0.7), node_w - Inches(0.4), Inches(1.0),
         desc, sz=14, color=DARK)

    # Arrow between nodes
    if i < len(nodes) - 1:
        ax1 = nx + node_w + Inches(0.05)
        ax2 = nx + node_w + node_gap - Inches(0.05)
        _arrow_right(sl, ax1, node_y + node_h // 2, ax2, node_y + node_h // 2, nc, 2)

_txt(sl, Inches(0.8), Inches(5.5), Inches(11.5), Inches(0.4),
     "Open-source Python package \u2014 seeking broader adoption and collaborators",
     sz=16, color=MID, align=PP_ALIGN.CENTER)


# =====================================================================
# SLIDE 17 — Summary + Questions
# =====================================================================
sl = _blank()

_txt(sl, Inches(1.5), Inches(1.5), Inches(10), Inches(0.7),
     "Summary", sz=32, bold=True, color=DARK, align=PP_ALIGN.CENTER)

summaries = [
    ("muDM \u2014 one vector data model for visualization, ML, and analysis",
     ACCENT1),
    ("Complements OME-Zarr for the annotation layer",
     ACCENT3),
    ("Open source, production-tested, growing",
     ACCENT4),
]
for i, (text, sc) in enumerate(summaries):
    y = Inches(2.8) + i * Inches(0.9)
    _circle(sl, Inches(2.5), y + Pt(8), Inches(0.2), sc)
    _txt(sl, Inches(3.0), y, Inches(7.5), Inches(0.5),
         text, sz=20, color=DARK)

_txt(sl, Inches(1.5), Inches(5.8), Inches(10), Inches(0.8),
     "Questions?", sz=40, bold=True, color=ACCENT1, align=PP_ALIGN.CENTER)


# ── Save ─────────────────────────────────────────────────────────────
out_path = "/data/code/work/microjson/muDM_presentation.pptx"
prs.save(out_path)
print(f"Saved: {out_path}")
print(f"Slides: {len(prs.slides)}")
